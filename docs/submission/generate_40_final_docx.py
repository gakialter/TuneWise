from __future__ import annotations

import hashlib
import os
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SUBMISSION = ROOT / "docs" / "submission"
TEMPLATE = SUBMISSION / "template" / "【40强赛】企业名称🤝xxx队名｜参赛方案标题.docx"
OUTPUT = SUBMISSION / "TuneWise_40强_最终参赛方案_飞书导入版.docx"
ASSETS = SUBMISSION / "assets" / "40-final"
EXPECTED_INLINE_IMAGES = 8
EXPECTED_MIN_TABLES = 9
EXPECTED_FIGURE_TITLES = [
    "图 1｜传统调机流程 vs TuneWise 决策支持流程",
    "图 2｜TuneWise 双层 AI 架构",
    "图 3｜TuneWise 核心决策链",
    "图 4｜安全调参决策流程",
    "图 5｜Fixed Demo 真实运行界面",
    "图 6｜结合调机步骤的决策演示",
    "图 7｜飞书 Aily 工程解释｜Workflow + 真实问答",
    "图 8｜验证与证据总结",
]

TITLE = (
    "【40强赛】舜宇光学科技🤝工智跃迁｜"
    "TuneWise——AA 工站 AI 调机决策支持"
)
CHALLENGE = (
    "如果你是智造专家，你将如何借助 AI 设计并打造"
    "「智造调机助手」，提升现场良率？"
)
SUMMARY = (
    "TuneWise 是面向精密光学 Active Alignment（AA）工站的 Human-in-the-loop AI 调机决策支持原型。"
    "确定性核心给出异常分析、根因优先级、历史参考案例与调参候选，工程师审核确认后，"
    "先做执行前仿真验证，再演示本地模拟设备受控执行；飞书 Aily 仅负责知识检索与解释。"
)

BLUE = "2446A8"
BLUE_DARK = "10233F"
BLUE_SOFT = "EAF0FF"
PURPLE = "6F5CE7"
PURPLE_SOFT = "F1EEFF"
GREEN = "148F5D"
GREEN_SOFT = "E8F6EF"
ORANGE = "D97706"
ORANGE_SOFT = "FFF2DE"
RED = "C73F46"
RED_SOFT = "FDEBEC"
INK = "172B4D"
MUTED = "5D6B82"
LINE = "CCD6E5"
PAPER = "FFFFFF"
FONT = "Microsoft YaHei"
BULLET_NUM_ID = "99"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require_inputs() -> None:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(
            "Official 40-final template missing from the repository source-of-truth path: "
            f"{TEMPLATE}"
        )
    required = [
        ASSETS / "01_before_after.png",
        ASSETS / "02_dual_layer_ai_architecture.png",
        ASSETS / "03_core_pipeline.png",
        ASSETS / "04_safe_parameter_decision.png",
        ASSETS / "05_replay_opcua_safety.png",
        ASSETS / "06_demo_evidence_card.png",
        ASSETS / "07_aily_rag_workflow.png",
        ASSETS / "08_validation_summary.png",
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required embedded image(s):\n" + "\n".join(missing))


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def set_east_asia_font(element, font_name: str = FONT) -> None:
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    font_name: str = FONT,
) -> None:
    run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    set_east_asia_font(run._element, font_name)


def add_qformat(style) -> None:
    if style._element.find(qn("w:qFormat")) is None:
        qformat = OxmlElement("w:qFormat")
        # w:qFormat must precede paragraph/run/table properties in CT_Style.
        # Appending it after w:pPr/w:rPr produces a package that Word may open
        # but fails strict OOXML schema validation.
        style._element.insert_element_before(
            qformat,
            "w:locked",
            "w:personal",
            "w:personalCompose",
            "w:personalReply",
            "w:rsid",
            "w:pPr",
            "w:rPr",
            "w:tblPr",
            "w:trPr",
            "w:tcPr",
        )


def set_outline_level(style, level: int) -> None:
    p_pr = style._element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def add_paragraph_style(
    doc: Document,
    name: str,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
    before: float = 0,
    after: float = 4,
    line_spacing: float = 1.4,
    keep_next: bool = False,
    outline_level: int | None = None,
    base_style=None,
):
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base_style is not None:
        style.base_style = base_style
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    set_east_asia_font(style._element)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    if outline_level is not None:
        set_outline_level(style, outline_level)
        add_qformat(style)
    return style


def configure_styles(doc: Document) -> None:
    normal = add_paragraph_style(
        doc,
        "Normal",
        size=10.5,
        after=4,
        line_spacing=1.42,
    )
    normal._element.set(qn("w:default"), "1")

    title = add_paragraph_style(
        doc,
        "Title",
        size=23.5,
        color=BLUE_DARK,
        bold=True,
        after=5,
        line_spacing=1.12,
        keep_next=True,
        base_style=normal,
    )
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    project = add_paragraph_style(
        doc,
        "Project Name",
        size=18,
        color=BLUE,
        bold=True,
        before=2,
        after=2,
        line_spacing=1.15,
        keep_next=True,
        base_style=normal,
    )
    project.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = add_paragraph_style(
        doc,
        "Subtitle",
        size=12.5,
        color=MUTED,
        after=10,
        line_spacing=1.2,
        keep_next=True,
        base_style=normal,
    )
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_paragraph_style(
        doc,
        "Heading 1",
        size=17,
        color=BLUE_DARK,
        bold=True,
        before=14,
        after=8,
        line_spacing=1.15,
        keep_next=True,
        outline_level=0,
        base_style=normal,
    )
    add_paragraph_style(
        doc,
        "Heading 2",
        size=14,
        color=BLUE,
        bold=True,
        before=11,
        after=6,
        line_spacing=1.2,
        keep_next=True,
        outline_level=1,
        base_style=normal,
    )
    add_paragraph_style(
        doc,
        "Heading 3",
        size=11.5,
        color=INK,
        bold=True,
        before=8,
        after=4,
        line_spacing=1.25,
        keep_next=True,
        outline_level=2,
        base_style=normal,
    )

    figure_title = add_paragraph_style(
        doc,
        "Figure Title",
        size=10.5,
        color=BLUE_DARK,
        bold=True,
        before=8,
        after=4,
        line_spacing=1.2,
        keep_next=True,
        base_style=normal,
    )
    figure_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption = add_paragraph_style(
        doc,
        "Caption",
        size=9,
        color=MUTED,
        after=7,
        line_spacing=1.25,
        base_style=normal,
    )
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    small = add_paragraph_style(
        doc,
        "Small Text",
        size=9,
        color=MUTED,
        after=3,
        line_spacing=1.25,
        base_style=normal,
    )
    small.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def configure_bullet_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part._element
    # The official competition template contains single-level bullet
    # definitions whose ``w:lvl`` nodes omit the schema-required ``w:ilvl``
    # attribute.  Preserve those definitions, but normalize them before adding
    # TuneWise's own list so the generated package remains strict OOXML-valid.
    for existing_abstract in numbering.findall(qn("w:abstractNum")):
        for level_index, existing_level in enumerate(
            existing_abstract.findall(qn("w:lvl"))
        ):
            if existing_level.get(qn("w:ilvl")) is None:
                existing_level.set(qn("w:ilvl"), str(level_index))

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), "900")

    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "bullet")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.extend([start, num_format, level_text, level_justification])

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "390")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "390")
    indent.set(qn("w:hanging"), "210")
    p_pr.extend([tabs, indent])
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Symbol")
    r_fonts.set(qn("w:hAnsi"), "Symbol")
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    level.append(r_pr)
    abstract_num.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract_num)
    else:
        first_num.addprevious(abstract_num)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), BULLET_NUM_ID)
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "900")
    num.append(abstract_id)
    cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
    if cleanup is None:
        numbering.append(num)
    else:
        cleanup.addprevious(num)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.85)
    section.bottom_margin = Cm(1.85)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    pg_mar = section._sectPr.find(qn("w:pgMar"))
    if pg_mar is not None:
        pg_mar.set(qn("w:gutter"), "0")


def configure_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, size=8.5, color=MUTED)

    field_run = paragraph.add_run()
    set_run_font(field_run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instruction, separate, result, end])

    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def add_heading(
    doc: Document,
    text: str,
    level: int,
    *,
    page_break_before: bool = False,
):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break_before
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size={1: 17, 2: 14, 3: 11.5}[level],
        bold=True,
        color={1: BLUE_DARK, 2: BLUE, 3: INK}[level],
    )
    return paragraph


def add_body(
    doc: Document,
    text: str,
    *,
    first_line: bool = True,
    keep_next: bool = False,
    bold_prefix: str | None = None,
):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Pt(21) if first_line else None
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = False
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, size=10.5, bold=True, color=INK)
        tail = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(tail, size=10.5, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK)
    return paragraph


def add_bullet(doc: Document, text: str, *, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Cm(0.69)
    paragraph.paragraph_format.first_line_indent = Cm(-0.37)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.space_after = Pt(1.8)
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), BULLET_NUM_ID)
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.insert(0, num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, size=10.5, bold=True, color=INK)
        tail = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(tail, size=10.5, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK)
    return paragraph


def add_appendix_heading(doc: Document, text: str):
    paragraph = add_heading(doc, text, 2)
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_appendix_bullet(doc: Document, text: str):
    paragraph = add_bullet(doc, text)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0.8)
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.insert_element_before(
            shading,
            "w:noWrap",
            "w:tcMar",
            "w:textDirection",
            "w:tcFitText",
            "w:vAlign",
            "w:hideMark",
            "w:headers",
            "w:cellIns",
            "w:cellDel",
            "w:cellMerge",
            "w:tcPrChange",
        )
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.insert_element_before(
            tc_mar,
            "w:textDirection",
            "w:tcFitText",
            "w:vAlign",
            "w:hideMark",
            "w:headers",
            "w:cellIns",
            "w:cellDel",
            "w:cellMerge",
            "w:tcPrChange",
        )
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.insert_element_before(
            borders,
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
            "w:tblPrChange",
        )
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_width(table, width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.insert_element_before(
            layout,
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
            "w:tblPrChange",
        )
    layout.set(qn("w:type"), "fixed")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def fill_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = INK,
    size: float = 9.5,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.style = "Normal"
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.28
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_info_card(doc: Document) -> None:
    rows = [
        ("企业", "舜宇光学科技"),
        ("队名", "工智跃迁"),
        ("命题", CHALLENGE),
        ("一句话摘要", SUMMARY),
        (
            "成员介绍 & 分工",
            "钟秉辰｜个人参赛\n"
            "负责 TuneWise 的方案设计、机器学习与安全决策原型、执行前仿真验证与 OPC-UA 本地模拟环境验证、"
            "飞书 Aily 工程协作层集成，以及项目工程验证与参赛材料整理。",
        ),
        (
            "使用的飞书 AI 能力",
            "飞书 Aily：工作流应用；知识空间；知识检索（RAG）；"
            "大语言模型（LLM）；已发布对话应用。",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 17.15)
    set_table_borders(table, color="AEBFE2", size="7")
    header = table.rows[0]
    prevent_row_split(header)
    fill_cell(
        header.cells[0],
        "项目",
        bold=True,
        color=PAPER,
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    fill_cell(
        header.cells[1],
        "填写内容",
        bold=True,
        color=PAPER,
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_shading(header.cells[0], BLUE)
    set_cell_shading(header.cells[1], BLUE)
    set_cell_width(header.cells[0], 3.25)
    set_cell_width(header.cells[1], 13.90)
    repeat_header(header)
    for index, (label, value) in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        set_cell_width(row.cells[0], 3.25)
        set_cell_width(row.cells[1], 13.90)
        fill_cell(row.cells[0], label, bold=True, color=BLUE_DARK, size=9.5)
        fill_cell(row.cells[1], value, size=9.3)
        set_cell_shading(row.cells[0], BLUE_SOFT)
        set_cell_shading(row.cells[1], PAPER if index % 2 == 0 else "F8FAFE")


def add_simple_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 9.0,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, sum(widths))
    set_table_borders(table, color="BFCBE0", size="6")
    header = table.rows[0]
    prevent_row_split(header)
    repeat_header(header)
    for index, value in enumerate(headers):
        set_cell_width(header.cells[index], widths[index])
        fill_cell(
            header.cells[index],
            value,
            bold=True,
            color=PAPER,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(header.cells[index], BLUE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            set_cell_width(row.cells[col_index], widths[col_index])
            align = WD_ALIGN_PARAGRAPH.CENTER if col_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(
                row.cells[col_index],
                value,
                size=font_size,
                align=align,
            )
            set_cell_shading(row.cells[col_index], PAPER if row_index % 2 == 0 else "F7F9FC")
    tail = doc.add_paragraph(style="Small Text")
    tail.paragraph_format.space_after = Pt(1)


def add_callout(
    doc: Document,
    label: str,
    text: str,
    *,
    fill: str = BLUE_SOFT,
    border: str = "9BB0DE",
    label_color: str = BLUE_DARK,
) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 17.15)
    set_table_borders(table, color=border, size="7")
    row = table.rows[0]
    prevent_row_split(row)
    cell = row.cells[0]
    set_cell_width(cell, 17.15)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=100, start=150, bottom=100, end=150)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Normal"
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.3
    lead = paragraph.add_run(label + "：")
    set_run_font(lead, size=9.7, bold=True, color=label_color)
    body = paragraph.add_run(text)
    set_run_font(body, size=9.7, color=INK)
    spacer = doc.add_paragraph(style="Small Text")
    spacer.paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, title: str, path: Path, caption: str, *, width_cm=16.6) -> None:
    title_paragraph = doc.add_paragraph(style="Figure Title")
    title_paragraph.paragraph_format.keep_with_next = True
    title_paragraph.paragraph_format.keep_together = True
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=10.5, bold=True, color=BLUE_DARK)

    image_paragraph = doc.add_paragraph(style="Normal")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = None
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.keep_together = True
    run = image_paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Cm(width_cm))
    inline_shape._inline.docPr.set("name", title)
    inline_shape._inline.docPr.set("descr", title)

    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.paragraph_format.keep_together = True
    caption_run = caption_paragraph.add_run("图注：" + caption)
    set_run_font(caption_run, size=9, color=MUTED)


def build_document() -> Document:
    doc = Document(str(TEMPLATE))
    clear_template_body(doc)
    configure_styles(doc)
    configure_bullet_numbering(doc)
    configure_page(doc)
    configure_footer(doc)

    properties = doc.core_properties
    properties.title = TITLE
    properties.subject = "AI 先锋未来人才大赛 40 强最终参赛方案"
    properties.author = "工智跃迁"
    properties.keywords = "TuneWise, 飞书 Aily, Active Alignment, AI 调机决策支持"
    properties.comments = ""

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run(TITLE)
    set_run_font(title_run, size=23.5, bold=True, color=BLUE_DARK)

    project = doc.add_paragraph(style="Project Name")
    project.paragraph_format.keep_with_next = True
    project_run = project.add_run("TuneWise")
    set_run_font(project_run, size=18, bold=True, color=BLUE)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.keep_with_next = True
    subtitle_run = subtitle.add_run("Human-in-the-loop AI 调机决策支持原型")
    set_run_font(subtitle_run, size=12.5, color=MUTED)

    add_heading(doc, "一、参赛方案信息卡【必须包含】", 1)
    add_info_card(doc)

    add_heading(doc, "方案摘要", 2)
    add_body(
        doc,
        "TuneWise 聚焦精密光学主动对准（Active Alignment，AA）单一工站，是 Human-in-the-loop AI 调机决策支持原型。AI 完成异常分析、根因优先级、历史参考案例与调参候选；工程师审核并确认后，系统先做执行前仿真验证（Replay），再在全部安全门禁通过且显式启用 OPC-UA Sandbox 时演示本地模拟设备受控执行。大语言模型不生成新参数，也不进入设备控制链。",
    )
    add_body(
        doc,
        "已发布的飞书 Aily 应用承担生成式 AI 协作层。它检索 8 文件版本化知识包、固定 Demo 证据与调机过程 Demo 证据，将模型、规则、安全边界和结果组织为自然语言解释。两层 AI 通过版本化项目知识与固定证据协作，而非通过实时控制接口连接：TuneWise 确定性工业 AI 核心产生可验证决策证据，Aily 只负责检索、解释、追溯与问答。",
    )
    add_callout(
        doc,
        "事实边界",
        "本项目是比赛原型，不代表舜宇真实内部系统。固定 Demo、调机过程 Demo、模型开发资产与案例均为公开知识和规则约束的合成数据；normalized_score 不是校准故障概率；仿真验证结果只属于固定模拟环境；OPC-UA 只连接本地 loopback sandbox；Aily 无控制权限。项目未接真实光学产线，也未完成授权真实 AA 业务数据验证。",
        fill=RED_SOFT,
        border="E6AEB2",
        label_color=RED,
    )

    add_heading(doc, "二、 方案成果展示【必须包含】", 1)
    add_heading(doc, "1. 命题场景描述、问题描述及痛点说明", 2)
    add_body(
        doc,
        "精密光学 Active Alignment 需要同时观察中心与四角 MTF 等质量指标，并调整位置、倾角、焦向等多个参数。故障现象与根因并非一一对应：相似的四角下降可能来自平面倾斜、XY 偏心、平台不稳定、参考漂移或条件式 Z 离焦；同一参数变化也可能同时影响多个质量指标。因此，现场判断不能只看单个数值，更不能把一句自然语言建议直接当作可执行参数。",
    )
    add_body(
        doc,
        "本方案聚焦方法级痛点，不将行业共性误述为舜宇已确认的内部现状。TuneWise 面向以下六类相互关联的问题：",
    )
    for text, prefix in [
        ("多指标耦合：中心与四角 MTF、位置、倾角、焦向和平台状态必须放在同一上下文判断。", "多指标耦合："),
        ("根因非一一对应：同一症状可能对应多类根因，单一数值不足以形成可靠排查顺序。", "根因非一一对应："),
        ("经验依赖：分散的信息与个人经验难以复核，判断依据也不易交接。", "经验依赖："),
        ("案例适用条件：历史动作缺少产品、工站、版本和审核状态时，可能被错误复用。", "案例适用条件："),
        ("调参安全：范围、步长、方向、跨标称值和版本过期都需要统一门禁。", "调参安全："),
        ("全链追溯（Provenance）：从输入、模型、规则、候选、工程师确认到结果都需要可追溯凭证。", "全链追溯（Provenance）："),
    ]:
        add_bullet(doc, text, bold_prefix=prefix)
    add_body(
        doc,
        "传统路径通常是“异常出现 → 人工查看多个指标 → 依赖经验猜测根因 → 翻找历史记录或案例 → 尝试参数 → 重新测试 → 反复迭代 → 决策依据难以追溯”。TuneWise 不替代工程师，而是把“看什么、为什么、能否调整、由谁确认、如何验证”变成结构化且可拒绝的链路。",
    )
    add_figure(
        doc,
        "图 1｜传统调机流程 vs TuneWise 决策支持流程",
        ASSETS / "01_before_after.png",
        "该图是工作流能力对比，不代表已验证的真实效率或良率 KPI。",
    )

    add_heading(doc, "2. 方案优势/创新点", 2)
    add_heading(doc, "2.1 从直接生成调参值转向证据约束型决策", 3)
    add_body(
        doc,
        "TuneWise 不允许模型直接给出可执行参数。根因排序只是起点；调参候选还必须具备根因证据、方向证据、当前适用案例证据、确定性幅值规则与参数安全校验（ParameterSafetyValidator）结果。缺少支持、方向冲突、版本过期或内容被篡改时，系统拒绝继续。AI 的责任由“替人拍板”收敛为“帮助人形成可验证判断”。",
    )
    add_heading(doc, "2.2 安全关键工业 AI 与生成式 AI 解耦", 3)
    add_body(
        doc,
        "底层 TuneWise 确定性工业 AI 核心负责数值分析、历史参考案例检索、规则约束、工程师确认、仿真验证与执行门禁；上层飞书 Aily 生成式 AI 协作层负责检索、解释、追溯与问答。Aily 能解释为什么 PLANE_TILT 排在第一、为什么选择 pitch 减少 1 tick、仿真验证通过说明什么，却不能生成或修改参数、创建 ConfirmedPlan、触发仿真验证或写入 OPC-UA。",
    )
    add_heading(doc, "2.3 执行前仿真验证与独立设备门禁", 3)
    add_body(
        doc,
        "新参数不会从模型输出直接流向设备。候选先经安全校验，再由工程师确认并冻结为不可变计划；仿真验证先复现导入基线，再在相同隐藏场景、扰动和 seed 下只改变确认参数。设备执行是独立边界，要求仿真验证通过、独立工程师确认和全部门禁再次通过。",
    )
    add_heading(doc, "2.4 从数据到执行凭证的全链 provenance", 3)
    add_body(
        doc,
        "数据、规则、模型、标准化器（StandardScaler）、案例索引、方向证据、候选、ConfirmedPlan、仿真验证结果（ReplayResult）与设备执行凭证均绑定版本或 SHA-256。写入开始后如果结果不确定，系统进入 UNKNOWN_OUTCOME，只用同一幂等键做只读核对，不盲目重写，也不根据当前值猜测成功。",
    )

    add_heading(doc, "3. 具体方案说明（突出AI能力）", 2)
    add_heading(doc, "3.1 双层 AI 总体架构", 3)
    add_body(
        doc,
        "上层飞书 Aily 生成式 AI 协作层包含知识空间、知识检索（RAG）、大语言模型（LLM）、工程解释与评委 / 工程问答，负责检索、解释、追溯与问答。下层 TuneWise 确定性工业 AI 核心包含 SPC 异常检测、工程特征提取、AI 根因优先级、历史参考案例检索、安全调参候选生成、参数安全校验、工程师确认、执行前仿真验证与本地模拟设备受控执行。",
    )
    add_body(
        doc,
        "两层之间传递的是版本化项目知识、固定 Demo 证据与调机过程 Demo 证据，不是实时控制接口。Core 产出固定证据，Aily 只读检索并解释；职责隔离使生成式 AI 的可用性不会扩大参数与设备权限。",
    )
    add_figure(
        doc,
        "图 2｜TuneWise 双层 AI 架构",
        ASSETS / "02_dual_layer_ai_architecture.png",
        "Core 产生确定、可审计的决策证据，Aily 只读检索并解释这些证据；Aily 不生成新参数、不创建 ConfirmedPlan、不触发仿真验证，也没有 OPC-UA 写权限。",
    )

    add_heading(doc, "3.2 TuneWise Core 技术链路", 3)
    add_body(
        doc,
        "完整链路为：版本化 CSV → 数据校验 → SPC 异常检测 → 工程特征提取 → 标准化器（StandardScaler）→ 多分类逻辑回归 → AI 根因优先级 → 历史参考案例检索 → 安全调参候选生成 → 参数安全校验 → 工程师确认 → ConfirmedPlan → 执行前仿真验证 → 本地模拟设备受控执行。调机过程信息首先在“当前适用案例”环节起作用，不进入根因分类器。",
        first_line=False,
    )
    add_figure(
        doc,
        "图 3｜TuneWise 核心决策链",
        ASSETS / "03_core_pipeline.png",
        "固定工程特征驱动根因排序；当前异常与根因证据再结合调机过程信息，决定当前适用案例。调机过程信息不进入分类器，最终动作仍受参数安全校验与工程师确认约束。",
    )

    add_heading(doc, "3.2.1 版本化数据、SPC 与工程特征", 3)
    add_body(
        doc,
        "输入是固定 schema 的 AA 批次 CSV、Dataset Manifest、规则与版本快照。导入同时校验原始文件 SHA-256 和规范化观测 SHA-256；固定 Demo 包含 24 条观测。SPC 使用冻结控制限与持续性规则，把批次路由为 TARGET_ANOMALY、NORMAL、NON_TARGET_GLOBAL_DEGRADATION 或 INSUFFICIENT_DATA，避免单点噪声直接触发诊断。",
    )
    add_body(
        doc,
        "系统形成 50 个固定批次工程特征，包括中心与四角 MTF、五个参数、振动、重复定位误差、标定残差的均值、标准差与趋势，以及最差角、四角极差、四角标准差、左右差、上下差和对角差等空间特征。哈希证明内容与版本一致，不证明数据来自真实产线。",
    )

    add_heading(doc, "3.2.2 可解释机器学习根因排序", 3)
    add_body(
        doc,
        "固定 StandardScaler 与 multinomial logistic regression 对五类候选根因排序：PLANE_TILT、XY_DECENTER、PLATFORM_INSTABILITY、REFERENCE_DRIFT、Z_DEFOCUS_CONDITIONAL；Z 类还受额外硬门控。模型输出稳定 Top-3、raw logit、主要 logit contribution 与 normalized_score。",
    )
    add_callout(
        doc,
        "Score 事实边界",
        "normalized_score 是当前可见候选根因的 relative ranking score，不是经真实产线故障频率校准的 calibrated fault probability，也不是现场发生率或真实置信度。logit contribution 只表示标准化特征对当前类别 raw logit 的贡献，不代表因果贡献。",
    )

    add_heading(doc, "3.2.3 调机过程信息与 APPROVED-only 历史参考案例检索", 3)
    add_body(
        doc,
        "历史参考案例检索对当前任务重新计算 50 维查询特征，并核对产品、工站、schema、特征定义与检索规则兼容性；随后在独立标准化器下以标准化欧氏距离排序。索引中的 60 个 APPROVED 案例来自合成开发训练分区，每类根因 12 个。APPROVED 只表示通过项目内版本化索引准入，不代表专家 Ground Truth、企业内部案例或真实生产验证；历史动作也不会被直接复制为当前参数。",
    )
    add_body(
        doc,
        "系统新增“调机过程信息”，包括当前调机阶段、上一步调整和上一步调整结果。它与当前异常及根因证据共同确定当前适用案例，继而影响案例参考方案（CASE_GUIDED）的支持证据。调机过程信息不进入 Logistic Regression classifier，不修改 StandardScaler，不改变 Root Cause Top-3 算法，也不直接计算参数值；真实运行 A/B 画面集中在 3.6 节展示。",
    )

    add_heading(doc, "3.3 安全参数决策", 3)
    add_body(
        doc,
        "参数方向来自冻结工程规则，不来自 LLM，也不由案例投票决定。固定 Demo 中，PLANE_TILT 的 pitch 使用 top_bottom_difference；该特征为正，当前 pitch 高于标称值，因此安全方向为 DECREASE。roll 对应的 left_right_difference 低于最小空间支持阈值，状态为 INSUFFICIENT_SUPPORT，不生成 roll 候选。",
    )
    add_body(
        doc,
        "幅值有三种已实现策略：保守调整方案（CONSERVATIVE）固定 1 tick，标准调整方案（STANDARD）固定 2 ticks，案例参考方案（CASE_GUIDED）使用当前适用 APPROVED 案例中方向一致且历史安全通过动作的 tick 绝对值中位数。一个 tick 为 0.050000 normalized prototype unit。所有候选都经过同一个服务端参数安全校验器，检查范围、网格、非零变化、最大变化、参数族、方向、标称值、版本与 candidate hash。",
    )
    add_body(
        doc,
        "工程师只选择服务端已有且 PASSED 的候选，前端不能覆盖参数。服务端复核候选、规则与哈希后形成状态为 VALID 的不可变 ConfirmedPlan；上游内容变化会使其 STALE，阻断仿真验证与设备执行。",
    )
    add_figure(
        doc,
        "图 4｜安全调参决策流程",
        ASSETS / "04_safe_parameter_decision.png",
        "三个候选均通过统一参数安全校验，最终由工程师确认（AA_PROCESS_ENGINEER）选择保守调整方案 pitch -1 tick；该方案使用固定规则幅值，不依赖案例幅值。证据不足的 roll 轴没有候选。",
    )

    add_heading(doc, "3.4 执行前仿真验证与 OPC-UA 工业安全门禁", 3)
    add_body(
        doc,
        "安全链为：已确认调参方案 → 状态检查 → 安全复核 → 基线复现 → 调参方案仿真 → 仿真验证结果 → 设备执行资格检查 → 本地模拟设备执行。执行前仿真验证只接受服务端保存且仍为 VALID 的 ConfirmedPlan；基线与调参方案共享 simulator、固定 seed、固定 disturbance、样本数、时间序列、schema 与版本，唯一允许变化的是已确认参数。",
        first_line=False,
    )
    add_body(
        doc,
        "设备执行默认关闭，只能在显式 OPCUA_SANDBOX 模式下连接固定 loopback endpoint，并重新检查计划 freshness、仿真验证绑定、参数安全、设备身份、映射、单位、访问能力与设备健康。仿真验证通过只解锁执行资格检查，不会自动写入；写入结果不确定时进入 UNKNOWN_OUTCOME，只按同一幂等键做只读核对，不盲目重试。",
    )
    add_callout(
        doc,
        "仿真验证 / OPC-UA 事实边界",
        "仿真验证通过只表示当前固定 simulator、seed、disturbance 和评价规则下，确认方案满足预设条件；不代表真实设备有效、良率提升、参数最优、生产收益或真实因果证明。当前 OPC-UA 仅为 127.0.0.1 loopback 的 LOCAL_OPCUA_SANDBOX；项目未接真实光学设备或生产线，也未完成真实 PLC 联锁、证书身份或现场安全认证。",
        fill=ORANGE_SOFT,
        border="E6BE83",
        label_color=ORANGE,
    )

    add_heading(doc, "3.5 固定 Demo 完整证据链", 3)
    add_body(
        doc,
        "固定 Demo 身份为 Task tw-demo-task-001，预置资产 tw-aa-demo-v1，Batch tw-aa-demo-batch-001，Station AA，产品型号 TW-AA-PROTOTYPE-V1，共 24 条规则约束的合成 AA 原型观测。下表事实由当前仓库资产与完整 API 链再次核实。",
    )
    add_simple_table(
        doc,
        ["环节", "固定证据", "事实边界"],
        [
            ["异常检测", "TARGET_ANOMALY；24/24 持续违反路由信号", "版本化 SPC，不是真实产线故障结论"],
            ["诊断", "PLANE_TILT；normalized_score 0.997781", "相对排序分数，不是 99.7781% 故障概率"],
            ["根因 Top-3", "PLANE_TILT → XY_DECENTER → REFERENCE_DRIFT", "候选优先级，不证明真实因果"],
            ["调参候选", "pitch 0.250000 → 0.200000；保守调整方案（CONSERVATIVE）；-1 tick", "固定规则候选，不是最优参数"],
            ["安全校验", "Candidate PASSED；roll INSUFFICIENT_SUPPORT", "只保留具备方向证据的单轴候选"],
            ["工程师确认", "AA_PROCESS_ENGINEER；ConfirmedPlan VALID", "固定演示身份，不映射真实企业授权体系"],
            ["仿真验证", "基线复现 PASSED；SUCCESS；attempt 1", "固定版本确定性模拟环境"],
            ["受控执行", "LOCAL_OPCUA_SANDBOX；SUCCEEDED；readback 0.200000", "本地 loopback 模拟设备，不是真实设备"],
        ],
        [2.25, 8.0, 6.9],
        font_size=8.5,
    )
    add_body(
        doc,
        "同一规划结果还生成标准调整方案（STANDARD：pitch 0.250000 → 0.150000，-2 ticks）与案例参考方案（CASE_GUIDED：pitch 0.250000 → 0.100000，-3 ticks），二者均 PASSED，但本次 ConfirmedPlan 仅选择保守调整方案。保守方案的幅值来自固定 1 tick 规则，不依赖案例幅值；案例参考方案的 -3 ticks 才使用当前适用 APPROVED 案例的中位数证据。",
    )
    add_callout(
        doc,
        "固定模拟指标（非生产 KPI）",
        "基线 → 调参方案仿真：中心 MTF 均值 0.831003 → 0.832128；最差角点 MTF 0.567683 → 0.651478；四角极差 0.184837 → 0.102292；四角标准差 0.071709 → 0.042654；控制限检查 false → true；目标异常触发 true → false。以上只属于当前固定模拟条件。",
        fill=BLUE_SOFT,
        border="B8C7F0",
        label_color=BLUE_DARK,
    )
    add_figure(
        doc,
        "图 5｜Fixed Demo 真实运行界面",
        ASSETS / "05_replay_opcua_safety.png",
        "tw-demo-task-001 从异常诊断、调参候选、工程师确认、执行前仿真验证到本地模拟设备执行的真实运行画面。软件真实运行不等于真实生产验证：当前证据使用 synthetic fixed dataset 与 LOCAL_OPCUA_SANDBOX，不代表真实生产设备。",
    )

    add_heading(doc, "3.6 结合调机步骤的决策演示", 3)
    add_body(
        doc,
        "真实运行的 /process-aware-demo/ A/B 页面使用相同测量数据与相同根因证据，Top-1 均为 PLANE_TILT。A｜初始评估没有上一步调整，当前适用案例为 tw-aa-approved-011，案例参考方案为 pitch -3 ticks；B｜调整后评估记录 pitch 0.250000 → 0.200000 且未观察到显著改善（NO_MATERIAL_IMPROVEMENT），当前适用案例变为 tw-aa-approved-003，案例参考方案变为 pitch -4 ticks。两边的保守调整方案均为 -1 tick、标准调整方案均为 -2 ticks，参数安全校验均为 PASSED。",
    )
    add_figure(
        doc,
        "图 6｜结合调机步骤的决策演示",
        ASSETS / "06_demo_evidence_card.png",
        "两个场景使用相同测量与根因证据；TuneWise 根据不同调机过程信息选择不同的当前适用案例，因此案例参考方案不同。画面来自真实运行软件，但数据是 synthetic process-context fixtures；它只证明确定性的上下文敏感案例选择，不代表舜宇真实 SOP，也不证明真实生产调参准确率。",
    )

    add_heading(doc, "3.7 飞书 Aily 工程解释工作流（RAG）", 3)
    add_body(
        doc,
        "TuneWise Core 的输出包含模型版本、规则、哈希、候选状态、仿真验证检查与设备凭证。Aily 的价值不是替代 Core，而是把分散在项目知识、工程规则、固定 Demo 与调机过程 Demo 证据中的事实变成可对话的解释入口。",
    )
    for text, prefix in [
        ("应用：TuneWise；已发布对话应用。", "应用："),
        ("工作流：TuneWise Engineering Copilot；Start → Knowledge Space Retrieval → LLM → End。", "工作流："),
        ("知识空间：TuneWise Engineering Knowledge；在线知识包（Live Knowledge Pack）：8 / 8。", "知识空间："),
        ("知识检索（RAG）：Top K = 5；Threshold Filter = off。", "知识检索（RAG）："),
        ("大语言模型（LLM）：发布时配置为 Doubao-seed-2.0-Pro。", "大语言模型（LLM）："),
    ]:
        add_bullet(doc, text, bold_prefix=prefix)
    add_body(
        doc,
        "Aily 已在发布应用中完成人工界面 / 对话验收（manual UI / conversational acceptance validation）：在线知识包（Live Knowledge Pack）8 / 8、既有安全硬门禁（Legacy Safety Hard Gates）PASS、调机过程问答（Process-aware QA）PASS、已发布环境（Published Environment）PASS。这些是人工验收记录，不是自动化 benchmark，不代表模型准确率 100%，也不是生产环境验证。",
    )
    add_body(
        doc,
        "图 7 中的真实 Aily 对话截图是固定 Demo 工程问答，不冒充调机过程 Q5 的 Live 截图；Process-aware QA PASS 来自单独完成的发布环境人工对话验收记录。",
    )
    add_callout(
        doc,
        "Aily 事实边界",
        "Aily 只负责知识检索、工程解释、证据追溯与问答，不实时读取 TuneWise 运行状态，不生成或选择参数，不创建或修改 ConfirmedPlan，不触发仿真验证、本地模拟设备执行或 OPC-UA write。Aily 没有控制权限，工程师确认始终属于 TuneWise Core 授权链。",
        fill=PURPLE_SOFT,
        border="C3B8EF",
        label_color=PURPLE,
    )
    add_figure(
        doc,
        "图 7｜飞书 Aily 工程解释｜Workflow + 真实问答",
        ASSETS / "07_aily_rag_workflow.png",
        "左侧为中文版 Workflow，右侧为已人工确认的固定 Demo 真实问答截图；它不是调机过程 Q5 的 Live 截图。Live Knowledge Pack 8 / 8、Legacy Safety Hard Gates、Process-aware QA 与 Published Environment 状态均来自 manual UI / conversational acceptance validation。真实 Aily 截图只证明检索与解释能力，不是设备控制证据；Aily 无参数生成、仿真触发或 OPC-UA 写权限。",
    )

    add_heading(doc, "3.8 公开真实数据验证边界审查", 3)
    add_body(
        doc,
        "按照 Coach 建议，我们进一步核验了 Rikkyo University 发布的公开真实光学实验数据 LOROS。该数据属于 slanted-edge 光学测量的真实实验采集，包含可追溯的 MTF、SFR 与处理后 ROI，许可为 CC-BY-4.0；当前正式 record DOI 为 10.5281/zenodo.17493261，配套论文 DOI 为 10.1186/s40645-025-00783-7。它可以证明公开真实光学实验数据存在，以及 provenance、license 与 MTF / SFR 链路可核验。",
    )
    add_body(
        doc,
        "但 LOROS 不是 AA production data，不能提供 AA 生产身份、x / y / z / pitch / roll 设备状态、调机动作、根因 Ground Truth、参数方向、干预前后结果、调机阶段或生产产出，因此不能支持 TuneWise 当前诊断与荐参合同的真实业务效果验证。",
    )
    add_callout(
        doc,
        "语义映射审查结论：NO-GO",
        "为避免制造虚假的生产验证，TuneWise 未强行接入 LOROS，也未伪造 PLANE_TILT Ground Truth、把 edge angle 当 pitch、把 Pos 0..5 当空间五点、把 wavelength 当参数、复制中心 MTF 到四角或修改 frozen classifier contract。真实 AA 业务效果继续保留至授权企业数据验证阶段。",
        fill=ORANGE_SOFT,
        border="E6BE83",
        label_color=ORANGE,
    )

    add_heading(doc, "4. 方案价值", 2)
    add_heading(doc, "4.1 工程效率：减少证据整理路径", 3)
    add_body(
        doc,
        "TuneWise 把异常指标、根因候选、案例、参数依据、安全检查、确认与结果放在同一任务上下文中，使讨论围绕同一版本证据展开；Aily 进一步把这些证据转成可检索的自然语言解释。实际节省时间仍需在企业只读影子模式（Shadow Mode）中测量。",
    )
    add_heading(doc, "4.2 调机安全：把“不能继续”变成系统能力", 3)
    add_body(
        doc,
        "数据不足、不可调根因、方向冲突、越界、离网格、跨标称、计划过期、仿真验证不通过、设备身份或能力不匹配都会阻断后续。设备执行默认关闭，当前只验证本地 sandbox 单参数路径，不能替代真实设备的 PLC 联锁、证书、审批与现场安全制度。",
    )
    add_heading(doc, "4.3 知识沉淀：保存证据，而不只保存结论", 3)
    add_body(
        doc,
        "可复用知识应包含异常现象、数据版本、模型与规则、根因排序、调机过程信息、候选依据、确认主体、仿真验证结果与适用条件。当前 APPROVED 案例库是固定合成开发资产；其准入隔离和版本化结构体现“知识必须先审核再复用”的原则。",
    )
    add_heading(doc, "4.4 可复制性：场景知识与安全骨架分离", 3)
    add_body(
        doc,
        "根因、特征、方向规则与约束属于具体工艺；版本、哈希、确认、仿真验证、幂等、receipt 与事实边界属于可复用安全骨架。迁移时可以替换场景资产，同时保留证据链和执行门禁思路。",
    )
    add_heading(doc, "4.5 价值验证：先定义量尺，再讨论收益", 3)
    add_body(
        doc,
        "进入企业只读影子模式（Shadow Mode）后，工程效率可记录从异常出现到形成完整证据包的耗时、人工检索案例步骤数和证据可理解性评价；安全性可记录 Validator 拒绝的无效候选、过期计划和门禁不满足事件；知识价值可记录证据包完整率、案例准入审核结果与解释一致性；业务结果则由企业在获授权数据上定义质量指标、对照方式和观察周期。当前原型没有目标值，不预先承诺效率提升、成本节省、良率增长或收益金额。",
    )

    add_heading(doc, "5. 方案体验入口 & demo展示视频 【建议包含】", 2)
    add_body(
        doc,
        "项目代码与工程证据：https://github.com/gakialter/TuneWise",
        first_line=False,
        bold_prefix="项目代码与工程证据：",
    )
    add_body(
        doc,
        "飞书 Aily：TuneWise Engineering Copilot 已发布，核心工作流与真实问答效果见上文截图，并已完成发布环境下的人工 UI / conversational acceptance validation。",
        first_line=False,
        bold_prefix="飞书 Aily：",
    )
    add_body(
        doc,
        "Demo 展示：固定 Demo 的异常检测、根因优先级、安全调参候选、工程师确认、仿真验证与本地模拟设备执行结果已在上文通过完整证据链展示；调机过程 A/B Demo 与固定 Demo 明确分离。",
        first_line=False,
        bold_prefix="Demo 展示：",
    )

    add_heading(doc, "三、自由展示区【可选 · 加分项】", 1)
    add_figure(
        doc,
        "图 8｜验证与证据总结",
        ASSETS / "08_validation_summary.png",
        "当前工程自动化验证为 Backend 454 / 454 PASS、Frontend 46 / 46 PASS；固定 Demo 与调机过程 Demo 浏览器验证均 PASS。Aily 状态属于人工 UI / 对话验收，公开真实数据审查为 NO-GO，不是验证 PASS。",
    )
    add_heading(doc, "验证与证据分层", 2)
    for text, prefix in [
        ("工程自动化验证：Backend 454 / 454 PASS；Frontend 46 / 46 PASS。", "工程自动化验证："),
        ("Demo 浏览器验证：固定 Demo PASS；调机过程 Demo（Process-aware Demo）PASS；桌面 / 移动端 PASS。固定 Demo 的本地模拟设备执行仍为 SUCCEEDED，readback 0.200000。", "Demo 浏览器验证："),
        ("Aily 人工验收：在线知识包（Live Knowledge Pack）8 / 8；既有安全硬门禁（Legacy Safety Hard Gates）PASS；调机过程问答（Process-aware QA）PASS；已发布环境（Published Environment）PASS。验证类型为 manual UI / conversational acceptance validation，不是自动化 benchmark。", "Aily 人工验收："),
        ("公开真实数据审查：LOROS provenance、license、MTF / SFR / ROI 可核验，但与 AA 调机合同存在语义缺口，结论为 NO-GO。", "公开真实数据审查："),
        ("历史自动化快照：commit 1fdc526cc8794965b6c591a2fa5bc399e50a4e2b 上 433 backend tests、40 frontend tests 与 production build 通过；仅作历史证据，不作为当前主验证数字。", "历史自动化快照："),
    ]:
        add_bullet(doc, text, bold_prefix=prefix)

    add_heading(doc, "当前原型限制", 2)
    for text in [
        "尚未接入舜宇或其他真实光学产线、真实设备、MES 或 QMS。",
        "尚未完成经授权真实 AA 业务数据验证；当前固定 Demo、调机过程 Demo、模型开发资产与 APPROVED cases 均不是企业真实生产数据。",
        "公开 LOROS 数据审查为 NO-GO / semantic mismatch，不能被写成真实 AA 验证 PASS。",
        "尚未验证真实 PLC/设备联锁、生产证书身份、现场安全协议、真实工艺控制限与物理单位映射。",
        "不支持生产级多参数原子执行、自动回滚或无人值守调机。",
        "尚无经真实业务数据验证的效率、成本、良率或收益 KPI。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "三阶段落地路线", 2)
    add_heading(doc, "阶段 1｜只读影子模式（Shadow Mode）", 3)
    add_body(
        doc,
        "取得授权导出数据，由设备、工艺与数据责任方确认字段语义、单位、设备型号、batch/lot、来源与授权引用。数据只进入 SHADOW_READ_ONLY，先做完整性、质量与离线输出评估，不进入训练集、案例库或设备写入。",
    )
    add_heading(doc, "阶段 2｜工程师决策支持", 3)
    add_body(
        doc,
        "在真实工程师监督下，对诊断证据、候选合法性与解释质量进行盲评或回顾性验证，建立真实审核协议、拒绝标准与版本管理。系统仍只提供决策支持，由工程师在企业既有流程中执行。",
    )
    add_heading(doc, "阶段 3｜受监督的受控执行", 3)
    add_body(
        doc,
        "只有设备侧受控 Method、CAS 或 PLC/上位机联锁、身份与证书、审批、回滚、并发所有权、断线恢复和现场安全测试全部获批后，才讨论受监督的真实设备集成。当前项目没有到达该阶段，也不以无人值守自动调机为近期目标。",
    )

    add_heading(doc, "推广边界", 2)
    add_body(
        doc,
        "TuneWise 优先适用于多参数耦合但参数空间可明确约束、过程指标可测量、根因候选可结构化、历史案例有适用条件与审核状态、高风险动作保留工程师确认、结果可先在离线或影子环境评估的设备调优问题。方法可优先评估迁移到光学 AA、精密装调及部分参数型工艺优化场景；每次迁移都必须重新建立真实特征、控制限、单位、规则、模型、案例准入与设备安全协议，不能直接复用当前合成参数。",
    )

    add_heading(doc, "四、附录", 1)
    add_appendix_heading(doc, "A. 事实边界摘要")
    for text in [
        "TuneWise 是面向 AA 单工站、人在回路的 AI 调机决策支持原型，不代表舜宇真实内部系统。",
        "固定 Demo、调机过程 Demo、模型开发资产与 APPROVED cases 为公开知识和规则约束的合成数据，不是舜宇真实生产数据，也不是专家 Ground Truth。",
        "normalized_score 是 relative ranking score，不是 calibrated fault probability。",
        "仿真验证通过只表示当前固定 simulator、seed、disturbance 和评价规则下满足预设条件，不证明真实良率、生产收益、真实设备效果、真实因果关系或参数最优。",
        "当前 OPC-UA 通道仅连接本地 loopback sandbox，不代表真实 PLC/设备接入或生产安全认证。",
        "Aily 只负责检索、解释、追溯与问答，不生成参数、不修改 ConfirmedPlan、不触发仿真验证、本地模拟设备执行或 OPC-UA write；工程师确认属于 Core 授权链。",
        "LOROS 仅用于公开真实数据的语义边界审查，结论为 NO-GO，不等于 AA production data 或真实业务验证。",
        "当前 V1 未使用 Bridge、MCP、HTTP Connector、Custom Connector、Runtime API 或 Web SDK。",
    ]:
        add_appendix_bullet(doc, text)
    add_callout(
        doc,
        "最终事实边界与结论",
        "项目未接真实光学产线，未完成授权真实 AA 业务数据验证。score、仿真验证、OPC-UA 与 Aily 分别保持“相对排序”“固定模拟环境”“本地 sandbox”“只读检索与解释”的能力边界。TuneWise 的核心价值，是把工业 AI 从“给出答案”推进到“形成一条可验证、可拒绝、可追溯、有人授权的决策链”。",
        fill=RED_SOFT,
        border="E6AEB2",
        label_color=RED,
    )

    add_appendix_heading(doc, "B. 验证口径说明")
    add_appendix_bullet(doc, "当前工程自动化验证：Backend 454 / 454 PASS；Frontend 46 / 46 PASS。")
    add_appendix_bullet(doc, "固定 Demo 与调机过程 Demo（Process-aware Demo）浏览器验证均 PASS；只证明本地演示表面行为，不证明真机或生产收益。")
    add_appendix_bullet(doc, "历史 433 backend / 40 frontend 只绑定 commit 1fdc526，仅作历史证据。")
    add_appendix_bullet(doc, "Aily 的 8 / 8 与三项 PASS 是人工 UI / conversational acceptance validation，不是模型准确率或 automated benchmark。")
    add_appendix_bullet(doc, "LOROS 审查是公开真实数据的语义映射 NO-GO，不是外部验证 PASS。")

    add_appendix_heading(doc, "C. 必要参考资料")
    add_appendix_bullet(doc, "TuneWise 项目代码、版本化工程文档与验证记录：https://github.com/gakialter/TuneWise")
    add_appendix_bullet(doc, "飞书 Aily：以项目内已发布应用配置、8 文件知识包与人工验收记录为准。")
    add_appendix_bullet(doc, "LOROS 当前记录 DOI：10.5281/zenodo.17493261；配套论文 DOI：10.1186/s40645-025-00783-7；许可：CC-BY-4.0。")
    add_appendix_bullet(doc, "OPC-UA：仅采用 TuneWise 本地 sandbox 的受控 Method、幂等、readback 与 reconciliation 语义，不主张 OPC-UA 协议天然提供 compare-and-set。")

    return doc


REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def sanitize_package(source: Path, target: Path) -> None:
    excluded = {
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/commentsIds.xml",
        "word/commentsExtensible.xml",
        "word/vbaProject.bin",
    }
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            name = info.filename
            if name in excluded:
                continue
            data = zin.read(name)
            if name.endswith(".rels"):
                root = ET.fromstring(data)
                for relationship in list(root):
                    rel_type = relationship.attrib.get("Type", "").lower()
                    rel_target = relationship.attrib.get("Target", "").lower()
                    target_mode = relationship.attrib.get("TargetMode", "").lower()
                    if (
                        "comments" in rel_type
                        or "comments" in rel_target
                        or "vbaproject" in rel_target
                        or target_mode == "external"
                    ):
                        root.remove(relationship)
                ET.register_namespace("", REL_NS)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif name == "[Content_Types].xml":
                root = ET.fromstring(data)
                for override in list(root):
                    part_name = override.attrib.get("PartName", "").lower()
                    content_type = override.attrib.get("ContentType", "").lower()
                    if "comments" in part_name or "comments" in content_type or "vbaproject" in part_name:
                        root.remove(override)
                ET.register_namespace("", CT_NS)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(info, data)


def save_document(doc: Document) -> None:
    build_path = OUTPUT.with_name(OUTPUT.stem + ".building.docx")
    clean_path = OUTPUT.with_name(OUTPUT.stem + ".clean.docx")
    for path in (build_path, clean_path):
        if path.exists():
            path.unlink()
    try:
        doc.save(str(build_path))
        sanitize_package(build_path, clean_path)
        os.replace(clean_path, OUTPUT)
    finally:
        for path in (build_path, clean_path):
            if path.exists():
                path.unlink()


def validate_generated_document(doc: Document) -> tuple[int, int]:
    headings = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name.startswith("Heading ")
    ]
    heading_text = "\n".join(paragraph.text for paragraph in headings)
    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    body_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    full_text = body_text + "\n" + table_text
    figure_titles = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name == "Figure Title"
    ]

    if not doc.paragraphs or not headings:
        raise RuntimeError("Generated DOCX has no body paragraphs or heading hierarchy")
    if len(doc.inline_shapes) != EXPECTED_INLINE_IMAGES:
        raise RuntimeError(
            "Generated DOCX image count mismatch: "
            f"expected {EXPECTED_INLINE_IMAGES}, got {len(doc.inline_shapes)}"
        )
    if figure_titles != EXPECTED_FIGURE_TITLES:
        raise RuntimeError(
            "Generated DOCX figure sequence mismatch: "
            f"expected {EXPECTED_FIGURE_TITLES}, got {figure_titles}"
        )
    if len(doc.tables) < EXPECTED_MIN_TABLES:
        raise RuntimeError(
            "Generated DOCX is missing required information/evidence tables: "
            f"expected at least {EXPECTED_MIN_TABLES}, got {len(doc.tables)}"
        )
    if len(doc.sections) != 1 or " PAGE " not in doc.sections[0].footer._element.xml:
        raise RuntimeError("Generated DOCX is missing its centered footer page-number field")

    required_headings = [
        "一、参赛方案信息卡【必须包含】",
        "二、 方案成果展示【必须包含】",
        "1. 命题场景描述、问题描述及痛点说明",
        "2. 方案优势/创新点",
        "3. 具体方案说明（突出AI能力）",
        "4. 方案价值",
        "5. 方案体验入口 & demo展示视频 【建议包含】",
        "三、自由展示区【可选 · 加分项】",
        "四、附录",
    ]
    missing_headings = [value for value in required_headings if value not in heading_text]
    if missing_headings:
        raise RuntimeError(
            "Generated DOCX is missing official-template section heading(s): "
            + ", ".join(missing_headings)
        )

    required_markers = [
        "AA 工站 AI 调机决策支持",
        "Human-in-the-loop AI 调机决策支持原型",
        "执行前仿真验证（Replay）",
        "调机过程信息不进入 Logistic Regression classifier",
        "tw-aa-approved-011",
        "tw-aa-approved-003",
        "NO_MATERIAL_IMPROVEMENT",
        "Live Knowledge Pack",
        "8 / 8",
        "Legacy Safety Hard Gates",
        "Process-aware QA",
        "Published Environment",
        "Backend 454 / 454 PASS",
        "Frontend 46 / 46 PASS",
        "10.5281/zenodo.17493261",
        "10.1186/s40645-025-00783-7",
        "语义映射审查结论：NO-GO",
        "PLANE_TILT；normalized_score 0.997781",
        "pitch 0.250000 → 0.200000",
        "基线复现 PASSED；SUCCESS；attempt 1",
        "LOCAL_OPCUA_SANDBOX；SUCCEEDED；readback 0.200000",
        "synthetic fixed dataset",
        "synthetic process-context fixtures",
        "不冒充调机过程 Q5 的 Live 截图",
        "真实 Aily 截图只证明检索与解释能力，不是设备控制证据",
    ] + EXPECTED_FIGURE_TITLES
    missing_markers = [value for value in required_markers if value not in full_text]
    if missing_markers:
        raise RuntimeError(
            "Generated DOCX is missing required current-fact marker(s): "
            + ", ".join(missing_markers)
        )

    stale_markers = [
        "7" + "/7",
        "7" + " / 7",
        "7" + "-file",
        "7" + " 文件 Knowledge Pack",
        "7" + " 文件知识包",
        "7" + " 个版本化知识文件",
        "Replay" + "-first",
        "Traditional AA Before" + " vs TuneWise After",
        "TuneWise Core" + " Pipeline",
        "Safe Parameter" + " Decision Flow",
        "Replay" + " + OPC-UA Safety Gate",
        "Feishu Aily" + " RAG Workflow",
        "Validation & Evidence" + " Summary",
        "图 5｜执行前仿真验证与设备安全门禁",
        "图 6｜tw-demo-task-001 固定 Demo 证据卡",
        "图 7｜飞书 Aily 工程解释工作流",
        "飞书 Aily 已发布工作流与固定 Demo 证据问答",
    ]
    stale_hits = [value for value in stale_markers if value in full_text]
    if stale_hits:
        raise RuntimeError(
            "Generated DOCX still contains stale judge-facing term(s): "
            + ", ".join(stale_hits)
        )

    return len(headings), len(doc.tables)


def main() -> None:
    require_inputs()
    template_hash = sha256(TEMPLATE)
    # This reopen is the hard gate: generation never falls back to a blank document.
    probe = Document(str(TEMPLATE))
    if len(probe.sections) != 1:
        raise RuntimeError("Official template must contain exactly one section")
    del probe

    document = build_document()
    save_document(document)

    reopened = Document(str(OUTPUT))
    heading_count, table_count = validate_generated_document(reopened)
    print(f"TEMPLATE_SHA256={template_hash}")
    # The official filename contains an emoji that the default Windows GBK
    # console cannot encode. The SHA-256 above is the reproducible identity.
    print("TEMPLATE=repository-official-template")
    print(f"OUTPUT={OUTPUT}")
    print(f"OUTPUT_BYTES={OUTPUT.stat().st_size}")
    print(f"PARAGRAPHS={len(reopened.paragraphs)}")
    print(f"HEADINGS={heading_count}")
    print(f"TABLES={table_count}")
    print(f"INLINE_IMAGES={len(reopened.inline_shapes)}")


if __name__ == "__main__":
    main()
